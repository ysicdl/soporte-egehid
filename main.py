import io
import os
import sys
import uuid
import sqlite3
from contextlib import contextmanager
from datetime import datetime

import openpyxl
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse, RedirectResponse
from pydantic import BaseModel
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# -------------------------------------------------------------
# INICIALIZACIÓN DE LA APLICACIÓN FASTAPI
# -------------------------------------------------------------
app = FastAPI(title="Sistema de Soporte EGEHID")

# En despliegue web el frontend se sirve desde el mismo origen, así que no se
# necesita CORS abierto. Se permite configurar orígenes extra por variable de
# entorno (separados por coma). Sin credenciales -> evita el combo inseguro
# "allow_origins=* + allow_credentials=True".
_origins = os.environ.get("ALLOWED_ORIGINS", "").strip()
allow_origins = [o.strip() for o in _origins.split(",") if o.strip()] if _origins else []
if allow_origins:
    app.add_middleware(
        CORSMiddleware,
        allow_origins=allow_origins,
        allow_credentials=False,
        allow_methods=["*"],
        allow_headers=["*"],
    )

# RUTA ABSOLUTA UNIFICADA DE LA BASE DE DATOS
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
# Permite reubicar la BD fuera del directorio de la app (recomendado en IIS,
# p. ej. una carpeta de datos con permisos de escritura para el App Pool).
DB_PATH = os.environ.get("SOPORTE_DB_PATH", os.path.join(BASE_DIR, "soporte.db"))


def obtener_ruta_recurso(nombre_archivo):
    """ Resuelve la ruta absoluta del archivo estático (HTML, ico, ...). """
    base_path = getattr(sys, "_MEIPASS", BASE_DIR)
    return os.path.join(base_path, nombre_archivo)


# -------------------------------------------------------------
# BASE DE DATOS SQLITE (CON MODO WAL)
# -------------------------------------------------------------
@contextmanager
def get_db():
    """ Conexión con cierre garantizado, incluso ante excepciones. """
    conn = sqlite3.connect(DB_PATH, timeout=10)
    try:
        yield conn
    finally:
        conn.close()


def init_db():
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("PRAGMA journal_mode=WAL;")
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS tickets (
                id_ticket TEXT PRIMARY KEY,
                tecnico TEXT,
                tecnico_companero TEXT,
                ubicacion TEXT,
                departamento TEXT,
                detalle_soporte TEXT,
                equipo TEXT,
                colaborador TEXT,
                nivel TEXT,
                estado TEXT,
                fecha TEXT,
                hora_inicio TEXT,
                hora_final TEXT,
                tiempo_total REAL,
                extension TEXT
            )
        ''')
        conn.commit()


init_db()

# -------------------------------------------------------------
# ENDPOINTS PARA SERVIR VISTAS HTML
# -------------------------------------------------------------
@app.get("/")
async def raiz():
    return RedirectResponse(url="monitoreo")


@app.get("/formulario")
async def servir_formulario():
    ruta = obtener_ruta_recurso("forms_soporte_ui.html")
    if not os.path.exists(ruta):
        raise HTTPException(status_code=404, detail="Vista de formulario no encontrada")
    return FileResponse(ruta)


@app.get("/monitoreo")
async def servir_monitoreo():
    ruta = obtener_ruta_recurso("monitoreo.html")
    if not os.path.exists(ruta):
        raise HTTPException(status_code=404, detail="Vista de monitoreo no encontrada")
    return FileResponse(ruta)


@app.get("/api/health")
async def health():
    return {"status": "ok"}


# -------------------------------------------------------------
# DATASET MAESTRO (EGEHID)
# -------------------------------------------------------------
TECNICOS = [
    "PAUL RAFAEL SANTANA HIERRO",
    "ESMERLIN ZACARIAS CUELLO",
    "MIGUEL ANGEL PEÑA",
    "CARLOS ANTONIO GRULLON",
    "OTRO"
]

UBICACIONES = [
    "NO APLICA", "EDIFICIO PRINCIPAL", "CENTRAL HIDROELÉCTRICA TAVERA",
    "CENTRAL HIDROELÉCTRICA ANGOSTURA", "CENTRAL HIDROELÉCTRICA MONCIÓN",
    "CENTRAL HIDROELÉCTRICA BAIGUAQUE", "CENTRAL HIDROELÉCTRICA BRAZO DERECHO",
    "CENTRAL HIDROELÉCTRICA INOA", "CENTRAL HIDROELÉCTRICA JIMENOA",
    "CENTRAL HIDROELÉCTRICA EL SALTO", "CENTRAL HIDROELÉCTRICA CONSTANZA",
    "CENTRAL HIDROELÉCTRICA RINCÓN", "CENTRAL HIDROELÉCTRICA HATILLO",
    "CENTRAL HIDROELÉCTRICA RÍO BLANCO", "CENTRAL HIDROELÉCTRICA JIGÜEY",
    "CENTRAL HIDROELÉCTRICA AGUACATE", "CENTRAL HIDROELÉCTRICA VALDESIA",
    "CENTRAL HIDROELÉCTRICA NIZAO", "CENTRAL HIDROELÉCTRICA NAJAYO",
    "CENTRAL HIDROELÉCTRICA PALOMINO", "CENTRAL HIDROELÉCTRICA SABANA YEGUA",
    "CENTRAL HIDROELÉCTRICA SABANETA", "CENTRAL HIDROELÉCTRICA LAS DAMAS",
    "CENTRAL HIDROELÉCTRICA LOS TOROS", "CENTRAL HIDROELÉCTRICA DOMINGO RODRÍGUEZ",
    "CENTRAL HIDROELÉCTRICA MAGUEYAL", "CENTRAL HIDROELÉCTRICA OCOA",
    "CENTRO DE TRANSPORTACIÓN QUITASUEÑO EN HAINA"
]

DEPARTAMENTOS = [
    "NO APLICA", "ALMACEN", "ARCHIVO GENERAL", "CAPACITACION", "CCH", "COMUNICACIONES",
    "CONSULTORIO MEDICO", "CONTRALORIA", "COOPERATIVA", "DDH", "DIGITALIZACION",
    "DIRECCION ADMINITRATIVO", "DIRECCION DE RRPP", "DIRECCION DE TECNOLOGIA",
    "DIRECCION FINANCIERO", "DIRECCION GESTION HUMANA", "DIRECCION JURIDICA",
    "GESTION DE CALIDAD", "LIBRE ACCESO A LA INFOMACION", "MANTENIMIENTO MECANICO",
    "OBRAS CIVILES", "PLANIFICACION", "PRENSA", "PROTOCOLO", "PROYECTO ESPECIAL",
    "RECEPCION", "RELACIONES INTERNACIONALES", "SECRETARIA GENERAL",
    "SERVICIOS GENERALES", "SEDE QUITA SUEÑO", "AUDITORIA", "BIENESTAR Y ASISTENCIA SOCIAL",
    "SALON DE REUNIONES 5TO. PISO", "GERENCIA MECANICA", "CONTABILIDAD",
    "DIRECCION DE MANTENIMIENTO", "SEGURIDAD MILITAR", "NOMINA", "CENTRAL TAVERA",
    "SALON DE CONSEJO 5TO. PISO", "COMPRAS", "5TO PISO", "SUB ADMINISTRACION TECNICA",
    "SUB DIRECCION", "ENERGIA RENOVABLE", "COMERCIALIZACIÓN", "PRESIDENCIA DEL CONSEJO", "ADMINISTRACION"
]

EQUIPOS = ["OTRO", "PC", "PRINTER", "PC/PRINTER", "PROYECTOR", "RED/INTERNET", "PROGRAMAS/APP", "PANTALLA/TV", "OTROS"]
NIVELES = ["NIVEL1", "NIVEL2", "NIVEL3"]
ESTADOS = ["RESUELTO", "NO RESUELTO", "PROCESO", "PROCESO/ RETIRO DE EQUIPO", "PENDIENTE", "ANULADO"]

# Conjuntos para validación (normalizados en mayúsculas)
_NIVELES_SET = {n.upper() for n in NIVELES}
_ESTADOS_SET = {e.upper() for e in ESTADOS}


# -------------------------------------------------------------
# MODELOS DE DATOS (PYDANTIC)
# -------------------------------------------------------------
class TicketRequest(BaseModel):
    tecnico: str
    ubicacion: str
    departamento: str
    detalle_soporte: str
    equipo: str
    colaborador: str
    nivel: str
    extension: str
    tecnico_companero: str = "NO APLICA"


class TicketEditarRequest(BaseModel):
    id_ticket: str
    tecnico: str
    tecnico_companero: str = "NO APLICA"
    colaborador: str
    extension: str
    ubicacion: str
    departamento: str
    equipo: str
    nivel: str
    detalle_soporte: str
    estado: str = "PENDIENTE"


class ActualizarEstadoRequest(BaseModel):
    id_ticket: str
    nuevo_estado: str


def _validar_nivel(nivel: str) -> str:
    n = (nivel or "").strip().upper()
    if n not in _NIVELES_SET:
        raise HTTPException(status_code=422, detail=f"Nivel inválido: {nivel}")
    return n


def _validar_estado(estado: str) -> str:
    e = (estado or "").strip().upper()
    if e not in _ESTADOS_SET:
        raise HTTPException(status_code=422, detail=f"Estado inválido: {estado}")
    return e


def _calcular_cierre(fecha_str, hora_inicio_str):
    """ Devuelve (hora_final, tiempo_total_min) al pasar a RESUELTO. """
    try:
        inicio_dt = datetime.strptime(f"{fecha_str} {hora_inicio_str}", "%Y-%m-%d %H:%M:%S")
        ahora_dt = datetime.now()
        diff_segundos = (ahora_dt - inicio_dt).total_seconds()
        return ahora_dt.strftime("%H:%M:%S"), round(max(diff_segundos, 0) / 60, 2)
    except Exception:
        return datetime.now().strftime("%H:%M:%S"), 1.0


# -------------------------------------------------------------
# ENDPOINTS API (SOPORTE Y BASE DE DATOS)
# -------------------------------------------------------------
@app.get("/api/catalogos")
async def obtener_catalogos():
    return {
        "tecnicos": TECNICOS,
        "ubicaciones": UBICACIONES,
        "departamentos": DEPARTAMENTOS,
        "equipos": EQUIPOS,
        "niveles": NIVELES,
        "estados": ESTADOS
    }


@app.post("/api/crear-ticket")
async def crear_ticket(ticket: TicketRequest):
    try:
        nivel = _validar_nivel(ticket.nivel)
        ahora = datetime.now()
        id_ticket = f"TK-{uuid.uuid4().hex[:6].upper()}"
        fecha = ahora.strftime("%Y-%m-%d")
        hora_inicio = ahora.strftime("%H:%M:%S")
        estado_inicial = "PENDIENTE"
        companero = ticket.tecnico_companero.strip() if ticket.tecnico_companero and ticket.tecnico_companero.strip() else "NO APLICA"

        with get_db() as conn:
            cursor = conn.cursor()
            cursor.execute('''
                INSERT INTO tickets (
                    id_ticket, tecnico, tecnico_companero, ubicacion, departamento,
                    detalle_soporte, equipo, colaborador, nivel, estado,
                    fecha, hora_inicio, hora_final, tiempo_total, extension
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                id_ticket, ticket.tecnico, companero, ticket.ubicacion, ticket.departamento,
                ticket.detalle_soporte, ticket.equipo, ticket.colaborador, nivel, estado_inicial,
                fecha, hora_inicio, "--", 0.0, ticket.extension.strip()
            ))
            conn.commit()

        return {"status": "success", "id_ticket": id_ticket, "hora_inicio": hora_inicio}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.put("/api/editar-ticket")
async def editar_ticket(ticket: TicketEditarRequest):
    try:
        nivel = _validar_nivel(ticket.nivel)
        nuevo_est = _validar_estado(ticket.estado)
        target_id = ticket.id_ticket.strip().upper()

        with get_db() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT fecha, hora_inicio FROM tickets WHERE UPPER(id_ticket) = ?", (target_id,))
            existente = cursor.fetchone()

            if not existente:
                raise HTTPException(status_code=404, detail=f"Ticket {ticket.id_ticket} no encontrado")

            hora_final = "--"
            tiempo_total = 0.0
            if nuevo_est == "RESUELTO":
                hora_final, tiempo_total = _calcular_cierre(existente[0], existente[1])

            cursor.execute('''
                UPDATE tickets
                SET tecnico = ?, tecnico_companero = ?, colaborador = ?, extension = ?,
                    ubicacion = ?, departamento = ?, equipo = ?, nivel = ?,
                    detalle_soporte = ?, estado = ?, hora_final = ?, tiempo_total = ?
                WHERE UPPER(id_ticket) = ?
            ''', (
                ticket.tecnico, ticket.tecnico_companero, ticket.colaborador, ticket.extension,
                ticket.ubicacion, ticket.departamento, ticket.equipo, nivel,
                ticket.detalle_soporte, nuevo_est, hora_final, tiempo_total, target_id
            ))
            conn.commit()

        return {"status": "success", "message": f"Ticket {target_id} actualizado exitosamente."}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.delete("/api/eliminar-ticket/{id_ticket}")
async def eliminar_ticket(id_ticket: str):
    try:
        target_id = id_ticket.strip().upper()
        with get_db() as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM tickets WHERE UPPER(id_ticket) = ?", (target_id,))
            filas_afectadas = cursor.rowcount
            conn.commit()

        if filas_afectadas == 0:
            raise HTTPException(status_code=404, detail=f"Ticket {id_ticket} no encontrado")

        return {"status": "success", "message": f"Ticket {target_id} eliminado exitosamente."}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/actualizar-estado")
async def actualizar_estado(req: ActualizarEstadoRequest):
    try:
        nuevo_est = _validar_estado(req.nuevo_estado)
        target_id = req.id_ticket.strip().upper()

        with get_db() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT fecha, hora_inicio FROM tickets WHERE UPPER(id_ticket) = ?", (target_id,))
            ticket = cursor.fetchone()

            if not ticket:
                raise HTTPException(status_code=404, detail=f"Ticket {req.id_ticket} no encontrado")

            hora_final = "--"
            tiempo_total = 0.0
            if nuevo_est == "RESUELTO":
                hora_final, tiempo_total = _calcular_cierre(ticket[0], ticket[1])

            cursor.execute('''
                UPDATE tickets
                SET estado = ?, hora_final = ?, tiempo_total = ?
                WHERE UPPER(id_ticket) = ?
            ''', (nuevo_est, hora_final, tiempo_total, target_id))
            conn.commit()

        return {"status": "success", "message": f"Estado de {req.id_ticket} actualizado a {nuevo_est}"}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/tickets")
async def obtener_tickets():
    try:
        with get_db() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT id_ticket, tecnico, tecnico_companero, ubicacion, departamento, detalle_soporte, equipo, colaborador, nivel, estado, fecha, hora_inicio, hora_final, tiempo_total, extension FROM tickets ORDER BY fecha DESC, hora_inicio DESC")
            rows = cursor.fetchall()

        tickets = []
        for row in rows:
            tickets.append({
                "id_ticket": row[0],
                "tecnico": row[1] or "",
                "tecnico_companero": row[2] or "NO APLICA",
                "ubicacion": row[3] or "",
                "departamento": row[4] or "",
                "detalle_soporte": row[5] or "",
                "equipo": row[6] or "",
                "colaborador": row[7] or "",
                "nivel": row[8] or "",
                "estado": row[9] or "PENDIENTE",
                "fecha": row[10] or "",
                "hora_inicio": row[11] or "",
                "hora_final": row[12] or "--",
                "tiempo_total": row[13] if row[13] is not None else 0.0,
                "extension": row[14] or "--"
            })
        return {"tickets": tickets}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/descargar-excel")
async def descargar_excel():
    try:
        with get_db() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT id_ticket, tecnico, tecnico_companero, ubicacion, departamento, detalle_soporte, equipo, colaborador, nivel, estado, fecha, hora_inicio, hora_final, tiempo_total, extension FROM tickets ORDER BY fecha DESC, hora_inicio DESC")
            filas = cursor.fetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reportes"

        headers = [
            "ID Ticket", "Técnico", "Técnico Compañero", "Ubicación", "Departamento",
            "Detalle Soporte", "Equipo", "Colaborador(a)", "Nivel", "Estado",
            "Fecha", "Hora Inicio", "Hora Final", "Tiempo Total (min)", "Extensión"
        ]
        ws.append(headers)

        for row in filas:
            ws.append(list(row))

        header_fill = PatternFill(start_color="0056D2", end_color="0056D2", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        column_widths = [14, 25, 25, 28, 25, 35, 15, 22, 12, 18, 14, 14, 14, 22, 12]
        for i, width in enumerate(column_widths, start=1):
            col_letter = openpyxl.utils.get_column_letter(i)
            ws.column_dimensions[col_letter].width = width

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        nombre_archivo = f"Reporte_Soportes_EGEHID_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# -------------------------------------------------------------
# ARRANQUE PARA DESARROLLO LOCAL
# -------------------------------------------------------------
# En IIS el proceso lo lanza HttpPlatformHandler (ver web.config) usando:
#   python -m uvicorn main:app --host 127.0.0.1 --port %HTTP_PLATFORM_PORT%
# Este bloque solo aplica al ejecutar "python main.py" en tu máquina.
if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("HTTP_PLATFORM_PORT", os.environ.get("PORT", "8000")))
    uvicorn.run(app, host="127.0.0.1", port=port, log_level="info")
